from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Luke_Tervit_VC_Pitch_Deck.pptx"
PROFILE = ROOT / "luke.png"
FIG_DIR = ROOT / "dissertation_figures" / "final"

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Avenir Next"

BG = RGBColor(247, 246, 242)
TEXT = RGBColor(24, 31, 41)
MUTED = RGBColor(95, 104, 117)
LINE = RGBColor(219, 223, 230)
NAVY = RGBColor(33, 55, 90)
BLUE = RGBColor(78, 120, 184)
GOLD = RGBColor(184, 129, 41)
SOFT = RGBColor(240, 242, 246)
WHITE = RGBColor(255, 255, 255)


def set_background(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def box(slide, x, y, w, h):
    return slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))


def rect(slide, x, y, w, h, fill, line=LINE, width=1.0):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(width)
    return shape


def rule(slide, x, y, w, color=LINE):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(0.015),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=16,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    shape = box(slide, x, y, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shape


def add_label(slide, text, x=0.85, y=0.58):
    add_text(slide, text.upper(), x, y, 3.0, 0.24, size=10, color=BLUE, bold=True)


def add_title(slide, text, x=0.85, y=1.0, w=7.6, h=1.1, size=27):
    add_text(slide, text, x, y, w, h, size=size, color=TEXT, bold=True)


def add_body(slide, text, x, y, w, h, size=15, color=MUTED):
    add_text(slide, text, x, y, w, h, size=size, color=color)


def add_bullets(slide, bullets, x, y, w, h, size=15, color=TEXT, bullet_color=GOLD, gap=10):
    shape = box(slide, x, y, w, h)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(gap)
        r1 = p.add_run()
        r1.text = "• "
        r1.font.name = FONT
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = bullet
        r2.font.name = FONT
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return shape


def add_metric(slide, value, label, x, y, w=2.2):
    add_text(slide, value, x, y, w, 0.5, size=28, color=NAVY, bold=True)
    add_text(slide, label, x, y + 0.48, w, 0.3, size=11, color=MUTED, bold=True)


def add_footer(slide, text):
    add_text(slide, text, 0.85, 7.03, 11.4, 0.18, size=8, color=MUTED)


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Luke Tervit | Founder")
    add_title(slide, "The digital wind tunnel\nfor human discourse", h=1.5, size=31)
    add_body(
        slide,
        "Pre-deployment simulation for high-stakes political and reputation-sensitive communication.",
        0.85,
        2.6,
        5.6,
        0.5,
        size=16,
    )
    rule(slide, 0.85, 3.35, 11.5)
    add_text(slide, "Overall accuracy distribution", 0.85, 3.62, 3.0, 0.2, size=12, color=BLUE, bold=True)
    slide.shapes.add_picture(str(FIG_DIR / "06_overall_accuracy_histogram.png"), Inches(0.85), Inches(3.95), width=Inches(6.15))
    add_body(
        slide,
        "Validated on 100 held-out threads using 0-shot forward simulation from the root post alone.",
        7.4,
        4.15,
        4.0,
        0.62,
        size=15,
        color=MUTED,
    )
    add_body(
        slide,
        "The system is strongest at forecasting conversation-level reaction patterns before they exist.",
        7.4,
        5.05,
        4.0,
        0.62,
        size=15,
        color=NAVY,
    )


def slide_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Problem")
    add_title(slide, "Current tools are autopsies,\nnot forecasts", h=1.35)
    add_body(
        slide,
        "Teams can measure what happened after launch. They still cannot test how the conversation is likely to unfold before posting.",
        0.85,
        2.45,
        6.1,
        0.62,
        size=16,
    )
    rect(slide, 0.85, 3.35, 5.55, 2.7, WHITE, line=LINE)
    rect(slide, 6.95, 3.35, 5.55, 2.7, WHITE, line=LINE)
    add_text(slide, "What teams have", 1.15, 3.7, 2.4, 0.25, size=12, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            "Social listening dashboards",
            "Engagement analytics",
            "Slow A/B tests and instinct",
        ],
        1.15,
        4.12,
        4.6,
        1.25,
        size=15,
    )
    add_text(slide, "What they need", 7.25, 3.7, 2.4, 0.25, size=12, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            "A pre-deployment stress test",
            "Variant comparison before exposure",
            "A view of the likely reply spiral",
        ],
        7.25,
        4.12,
        4.6,
        1.25,
        size=15,
    )
    add_body(
        slide,
        "The unit of risk is not the post. It is the conversation the post creates.",
        0.85,
        6.45,
        7.2,
        0.32,
        size=18,
        color=NAVY,
    )


def slide_mcdonalds(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Decision Failure Example")
    add_title(slide, "McDonald’s Big Arch shows the\nfailure mode clearly", h=1.3)
    add_body(
        slide,
        "The launch did not fail because the asset looked low quality. It failed because the internet read it differently than the internal team did.",
        0.85,
        2.45,
        6.5,
        0.62,
        size=16,
    )
    rect(slide, 0.85, 3.4, 5.45, 2.5, WHITE, line=LINE)
    rect(slide, 6.8, 3.4, 5.45, 2.5, WHITE, line=LINE)
    add_text(slide, "Internal read", 1.15, 3.75, 1.5, 0.22, size=12, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            "Polished and on-brand",
            "Safe executive launch content",
            "Looks professional in review",
        ],
        1.15,
        4.15,
        4.4,
        1.2,
        size=14,
    )
    add_text(slide, "Internet read", 7.1, 3.75, 1.5, 0.22, size=12, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            "Tiny bite and “product” language became the focal point",
            "The video read as overly corporate and inauthentic",
            "The backlash lived in the reply spiral, not the asset itself",
        ],
        7.1,
        4.15,
        4.45,
        1.45,
        size=14,
    )
    add_body(
        slide,
        "That is the category: not post-mortem analytics, but a pre-deployment simulation layer that flags likely reaction failure before launch.",
        0.85,
        6.32,
        10.4,
        0.42,
        size=15,
        color=NAVY,
    )


def slide_value_prop(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Value Prop")
    add_title(slide, "Simulate, compare, and choose\nbefore you publish", h=1.3)
    add_body(
        slide,
        "We do distributional forecasting. We predict the shape of the thread, not just a single engagement score.",
        0.85,
        2.45,
        6.4,
        0.58,
        size=16,
    )
    steps = [
        ("1", "Input a root post or several variants"),
        ("2", "Simulate historically grounded responders"),
        ("3", "Forecast sentiment, aggression, emotion, and political mix"),
        ("4", "Pick the strongest wording before launch"),
    ]
    y = 3.3
    for num, text in steps:
        rect(slide, 0.95, y, 11.3, 0.64, WHITE, line=LINE)
        add_text(slide, num, 1.2, y + 0.17, 0.3, 0.2, size=17, color=GOLD, bold=True)
        add_text(slide, text, 1.75, y + 0.15, 9.8, 0.24, size=16, color=TEXT, bold=True)
        y += 0.78
    add_body(
        slide,
        "This is not exact tweet prediction and not a polling replacement. It is bounded pre-deployment decision support.",
        0.85,
        6.62,
        9.0,
        0.28,
        size=14,
    )


def slide_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Underlying Magic")
    add_title(slide, "Historically grounded agents,\n0-shot forward simulation", h=1.3)
    add_body(
        slide,
        "The system is not a wrapper prompt. It is a simulation pipeline with grounded agents, synchronous rounds, and a shared evaluation stack.",
        0.85,
        2.45,
        6.5,
        0.62,
        size=16,
    )
    steps = [
        ("1", "Build responder personas", "Five classifiers build agent profiles from historical behavior."),
        ("2", "Start from the root post only", "Every run begins in 0-shot mode with no target replies exposed."),
        ("3", "Simulate the thread forward", "Agents interact over 10 synchronous rounds and condition on recent thread context."),
        ("4", "Score simulated vs real output", "The same evaluation stack measures how closely the simulated thread matches reality."),
    ]
    y = 3.4
    for num, title, body in steps:
        rect(slide, 0.95, y, 11.2, 0.68, WHITE, line=LINE)
        add_text(slide, num, 1.18, y + 0.2, 0.24, 0.2, size=16, color=GOLD, bold=True)
        add_text(slide, title, 1.7, y + 0.12, 2.65, 0.26, size=14, color=TEXT, bold=True)
        add_text(slide, body, 4.35, y + 0.14, 6.9, 0.24, size=13, color=MUTED)
        y += 0.84


def slide_business(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Business Plan")
    add_title(slide, "Start narrow, sell into a painful\npre-deployment workflow", h=1.3)
    rect(slide, 0.85, 2.55, 3.5, 3.25, WHITE, line=LINE)
    rect(slide, 4.7, 2.55, 3.5, 3.25, WHITE, line=LINE)
    rect(slide, 8.55, 2.55, 3.5, 3.25, WHITE, line=LINE)
    add_text(slide, "Initial buyer", 1.15, 2.9, 1.5, 0.22, size=12, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            "US campaigns",
            "PACs and consultants",
            "Digital strategy teams",
        ],
        1.15,
        3.3,
        2.6,
        1.15,
        size=15,
    )
    add_text(slide, "Product shape", 5.0, 2.9, 1.5, 0.22, size=12, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            "Message-testing workflow",
            "Variant comparison",
            "Managed analysis, then software",
        ],
        5.0,
        3.3,
        2.6,
        1.15,
        size=15,
    )
    add_text(slide, "Expansion", 8.85, 2.9, 1.5, 0.22, size=12, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            "Public affairs",
            "Reputation-sensitive enterprise",
            "API and agency workflows",
        ],
        8.85,
        3.3,
        2.75,
        1.15,
        size=15,
    )
    add_body(
        slide,
        "The wedge is politics because message risk is high, workflows are fast, and the pain is immediate.",
        0.85,
        6.35,
        8.0,
        0.3,
        size=16,
        color=NAVY,
    )


def slide_status(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Current Status")
    add_title(slide, "The research system is already\nvalidated on forward prediction", h=1.3)
    add_metric(slide, "92.3%", "mean overall accuracy on 100 held-out threads", 0.85, 2.55, 3.2)
    add_metric(slide, "94.54%", "strict chronological holdout weighted accuracy", 4.4, 2.55, 3.2)
    add_metric(slide, "97 / 100", "threads in the EXCELLENT band", 8.05, 2.55, 2.8)
    rule(slide, 0.85, 3.9, 11.4)
    add_bullets(
        slide,
        [
            "0-shot forward simulation from the root post alone",
            "No significant directional sentiment bias in the final pipeline",
            "21-configuration ablation study shows performance comes from system design, not prompt luck",
            "Working runtime, figures, evaluation assets, and pitchable proof already exist",
        ],
        0.85,
        4.3,
        7.2,
        1.6,
        size=15,
    )
    slide.shapes.add_picture(str(FIG_DIR / "06_overall_accuracy_histogram.png"), Inches(8.55), Inches(4.35), width=Inches(3.45))
    add_footer(slide, "Metrics sourced from dissertation.tex and whitepaper.tex.")


def slide_scaling(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "Scaling Plan")
    add_title(slide, "Turn validated simulation into a\ncompounding product loop", h=1.3)
    add_body(
        slide,
        "The moat is not model access. The moat is calibration from repeated real-world use.",
        0.85,
        2.45,
        5.6,
        0.42,
        size=16,
    )
    steps = [
        ("1", "Run simulations for real message decisions"),
        ("2", "Observe the actual launch outcome"),
        ("3", "Measure prediction gaps and recalibrate the system"),
        ("4", "Expand from politics into broader reputation workflows"),
    ]
    y = 3.18
    for num, text in steps:
        rect(slide, 1.0, y, 10.8, 0.62, SOFT, line=LINE)
        add_text(slide, num, 1.22, y + 0.16, 0.24, 0.2, size=16, color=GOLD, bold=True)
        add_text(slide, text, 1.7, y + 0.14, 9.6, 0.22, size=15, color=TEXT, bold=True)
        y += 0.76
    add_body(
        slide,
        "Simulations → outcomes → recalibration is the path from research system to defensible product.",
        0.85,
        6.72,
        8.6,
        0.28,
        size=16,
        color=NAVY,
    )


def slide_about(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_label(slide, "About Me")
    add_title(slide, "Technical founder with both\nresearch depth and product instinct", h=1.3)
    slide.shapes.add_picture(str(PROFILE), Inches(0.95), Inches(2.6), width=Inches(2.45))
    add_bullets(
        slide,
        [
            "Luke Tervit",
            "First Class BSc Computer Science, University of Edinburgh",
            "Product Engineering Intern at Granola AI",
            "Built the full research, runtime, and evaluation pipeline end to end",
            "Previously built and sold a viral hackathon startup within 20 hours",
        ],
        3.95,
        2.8,
        7.8,
        2.0,
        size=16,
    )
    add_body(
        slide,
        "This company needs someone who can bridge multi-agent systems, rigorous validation, and an operator-facing workflow. That is the founder fit.",
        3.95,
        5.55,
        7.7,
        0.45,
        size=16,
        color=NAVY,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide_title(prs)
    slide_problem(prs)
    slide_mcdonalds(prs)
    slide_value_prop(prs)
    slide_pipeline(prs)
    slide_business(prs)
    slide_status(prs)
    slide_scaling(prs)
    slide_about(prs)

    prs.save(OUT)


if __name__ == "__main__":
    build()
